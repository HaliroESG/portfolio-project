from __future__ import annotations

from openpyxl import Workbook

from scripts.import_target_allocations_excel import parse_target_excel, run_import


class _Response:
    def __init__(self, data=None):
        self.data = data or []


class _Table:
    def __init__(self, name: str, existing: set[tuple[str, str]]):
        self.name = name
        self.existing = existing
        self.operations = []
        self.filters = {}

    def select(self, columns):
        self.operations.append(("select", columns))
        return self

    def eq(self, column, value):
        self.filters[column] = value
        self.operations.append(("eq", column, value))
        return self

    def limit(self, value):
        self.operations.append(("limit", value))
        return self

    def update(self, payload):
        self.operations.append(("update", payload))
        return self

    def insert(self, payload):
        self.operations.append(("insert", payload))
        return self

    def execute(self):
        if self.name == "portfolio_positions" and self.operations and self.operations[0][0] == "select":
            key = (self.filters.get("portfolio_id"), self.filters.get("ticker"))
            return _Response([{"id": 1}] if key in self.existing else [])
        return _Response()


class _Supabase:
    def __init__(self, existing=None):
        self.existing = existing or set()
        self.tables = {}

    def table(self, name):
        table = _Table(name, self.existing)
        self.tables.setdefault(name, []).append(table)
        return table


def _write_target_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["portfolio_id", "ticker", "name", "asset_class", "currency", "target_weight_pct", "notes"])
    sheet.append(["main", "MSFT", "Microsoft", "Stock", "USD", 60, "core"])
    sheet.append(["main", "CW8", "Amundi MSCI World", "ETF", "EUR", 40, None])
    workbook.save(path)


def test_parse_target_excel_accepts_simple_allocation_file(tmp_path):
    source = tmp_path / "targets.xlsx"
    _write_target_workbook(source)

    report = parse_target_excel(source)

    assert report["rows_read"] == 2
    assert report["rejected"] == []
    assert [row.ticker for row in report["accepted"]] == ["MSFT", "CW8"]
    assert sum(row.target_weight_pct for row in report["accepted"]) == 100


def test_target_excel_import_apply_updates_existing_and_inserts_target_only(tmp_path):
    source = tmp_path / "targets.xlsx"
    _write_target_workbook(source)
    fake = _Supabase(existing={("main", "MSFT")})

    report = run_import(source, dry_run=False, supabase_client=fake)

    assert report["ok"] is True
    assert [row["ticker"] for row in report["updated"]] == ["MSFT"]
    assert [row["ticker"] for row in report["inserted"]] == ["CW8"]

    update_payload = next(
        op[1]
        for table in fake.tables["portfolio_positions"]
        for op in table.operations
        if op[0] == "update"
    )
    insert_payload = next(
        op[1]
        for table in fake.tables["portfolio_positions"]
        for op in table.operations
        if op[0] == "insert"
    )

    assert "quantity_current" not in update_payload
    assert "pru" not in update_payload
    assert update_payload["target_source"] == "excel"
    assert update_payload["target_source_file"] == "targets.xlsx"
    assert "target_updated_at" in update_payload
    assert insert_payload["quantity_current"] == 0
    assert insert_payload["target_weight_pct"] == 40


def test_target_excel_import_rejects_invalid_portfolio_total(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["portfolio_id", "ticker", "target_weight_pct"])
    sheet.append(["main", "MSFT", 80])
    source = tmp_path / "targets_bad.xlsx"
    workbook.save(source)

    report = run_import(source, dry_run=True)

    assert report["ok"] is False
    assert "must sum to 100%" in report["error"]
